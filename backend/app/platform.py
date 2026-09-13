from __future__ import annotations

import csv
import io
import json
import secrets
import string
import uuid
import re
from pathlib import Path

from app.auth import hash_password, verify_password
from app.database import UPLOAD_DIR, connect
from app.models import (
    AdminUserUpdate,
    Classroom,
    ClassroomCreate,
    ClassroomUpdate,
    DashboardStats,
    DiagnosisResult,
    ExerciseResult,
    ImportCommitResult,
    ImportPreview,
    KnowledgeGraph,
    KnowledgeNode,
    PasswordUpdate,
    ProfileUpdate,
    TeacherRegisterRequest,
    TeacherReview,
    Ticket,
    TicketCreate,
    TicketReplyCreate,
    TicketUpdate,
    User,
)


ORG_ID = "org_demo"


def initialize_extended_data() -> None:
    with connect() as connection:
        connection.execute(
            "INSERT OR IGNORE INTO organizations(id, name, code) VALUES (?, ?, ?)",
            (ORG_ID, "金扬智能示范学校", "JINYANG-DEMO"),
        )
        for name in ("计算机学院", "人工智能学院", "教务处"):
            connection.execute(
                "INSERT OR IGNORE INTO departments(id, organization_id, name) VALUES (?, ?, ?)",
                (f"dept_{uuid.uuid5(uuid.NAMESPACE_DNS, name).hex[:10]}", ORG_ID, name),
            )
        connection.execute("UPDATE users SET organization_id = ? WHERE organization_id = ''", (ORG_ID,))
        connection.execute("UPDATE courses SET organization_id = ? WHERE organization_id = ''", (ORG_ID,))
        _ensure_demo_profiles(connection)
        _ensure_graph_versions(connection)
        _ensure_default_classes(connection)
        _ensure_document_chunks(connection)


def _ensure_demo_profiles(connection) -> None:
    connection.execute(
        "UPDATE users SET department = '计算机学院', title = '讲师' WHERE id = 'teacher_demo' AND department = ''"
    )
    connection.execute(
        "UPDATE users SET student_no = '20260001' WHERE id = 'student_demo' AND student_no = ''"
    )


def _ensure_graph_versions(connection) -> None:
    courses = connection.execute("SELECT id, owner_id, active_version_id FROM courses").fetchall()
    for course in courses:
        if course["active_version_id"]:
            continue
        graph = _graph_dict(connection, course["id"])
        version_id = f"version_{uuid.uuid4().hex[:12]}"
        connection.execute(
            """
            INSERT INTO graph_versions(id, course_id, version_no, status, trigger, summary, graph_json, created_by)
            VALUES (?, ?, 1, 'active', 'migration', 'P0 图谱迁移', ?, ?)
            """,
            (version_id, course["id"], json.dumps(graph, ensure_ascii=False), course["owner_id"]),
        )
        connection.execute("UPDATE courses SET active_version_id = ? WHERE id = ?", (version_id, course["id"]))


def _ensure_default_classes(connection) -> None:
    courses = connection.execute("SELECT id, name, owner_id, organization_id FROM courses").fetchall()
    for index, course in enumerate(courses, start=1):
        class_id = f"class_{course['id']}"
        code = f"CG{index:04d}"
        connection.execute(
            """
            INSERT OR IGNORE INTO classrooms(
                id, organization_id, course_id, name, semester, join_code, status, primary_teacher_id
            ) VALUES (?, ?, ?, ?, '2026-2027 第一学期', ?, 'active', ?)
            """,
            (class_id, course["organization_id"] or ORG_ID, course["id"], f"{course['name']} 1班", code, course["owner_id"]),
        )
        connection.execute(
            """
            INSERT OR IGNORE INTO class_teachers(classroom_id, teacher_id, role, can_edit_course)
            VALUES (?, ?, 'primary', 1)
            """,
            (class_id, course["owner_id"]),
        )
        if connection.execute("SELECT 1 FROM users WHERE id = 'student_demo'").fetchone():
            connection.execute(
                "INSERT OR IGNORE INTO enrollments(classroom_id, student_id, status) VALUES (?, 'student_demo', 'active')",
                (class_id,),
            )
            connection.execute(
                """
                INSERT OR IGNORE INTO class_learning_progress(user_id, classroom_id, node_id, mastered, updated_at)
                SELECT user_id, ?, node_id, mastered, updated_at FROM learning_progress
                WHERE user_id = 'student_demo' AND course_id = ?
                """,
                (class_id, course["id"]),
            )


def _ensure_document_chunks(connection) -> None:
    documents = connection.execute(
        "SELECT id, course_id, content FROM documents WHERE status = 'active'"
    ).fetchall()
    for document in documents:
        if connection.execute("SELECT 1 FROM document_chunks WHERE document_id = ?", (document["id"],)).fetchone():
            continue
        _insert_chunks(connection, document["id"], document["course_id"], document["content"])
        source_count = connection.execute(
            "SELECT COUNT(*) FROM node_sources ns JOIN nodes n ON n.id = ns.node_id WHERE n.course_id = ?",
            (document["course_id"],),
        ).fetchone()[0]
        if source_count == 0:
            connection.execute(
                """
                INSERT OR IGNORE INTO node_sources(node_id, document_id, source_type, source_excerpt)
                SELECT id, ?, 'aigc', '' FROM nodes WHERE course_id = ?
                """,
                (document["id"], document["course_id"]),
            )
            connection.execute(
                """
                INSERT OR IGNORE INTO edge_sources(edge_id, document_id, source_type)
                SELECT id, ?, 'aigc' FROM edges WHERE course_id = ?
                """,
                (document["id"], document["course_id"]),
            )


def _insert_chunks(connection, document_id: str, course_id: str, content: str) -> None:
    from app.services.retrieval import index_text

    current_page: int | None = None
    page_by_paragraph: list[int | None] = []
    paragraphs: list[str] = []
    for item in content.replace("\r", "").split("\n\n"):
        paragraph = item.strip()
        if not paragraph:
            continue
        marker = re.match(r"^\[PAGE (\d+)\]\s*", paragraph)
        if marker:
            current_page = int(marker.group(1))
            paragraph = paragraph[marker.end():].strip()
        if paragraph:
            paragraphs.append(paragraph)
            page_by_paragraph.append(current_page)
    if not paragraphs:
        paragraphs = [content[:4000]] if content else []
        page_by_paragraph = [None] * len(paragraphs)
    chunks: list[tuple[str, int | None]] = []
    buffer = ""
    buffer_page: int | None = None
    for paragraph, page_no in zip(paragraphs, page_by_paragraph):
        if buffer and len(buffer) + len(paragraph) > 1800:
            chunks.append((buffer, buffer_page))
            buffer = paragraph
            buffer_page = page_no
        else:
            buffer = f"{buffer}\n\n{paragraph}".strip()
            if buffer_page is None:
                buffer_page = page_no
    if buffer:
        chunks.append((buffer, buffer_page))
    for index, (chunk, page_no) in enumerate(chunks):
        chunk_id = f"chunk_{uuid.uuid4().hex[:12]}"
        connection.execute(
            "INSERT INTO document_chunks(id, document_id, course_id, chunk_index, page_no, content) VALUES (?, ?, ?, ?, ?, ?)",
            (chunk_id, document_id, course_id, index, page_no, chunk),
        )
        try:
            connection.execute(
                "INSERT INTO document_chunks_fts(chunk_id, course_id, tokens) VALUES (?, ?, ?)",
                (chunk_id, course_id, index_text(chunk)),
            )
        except Exception:
            pass


def register_teacher(payload: TeacherRegisterRequest) -> User:
    username = payload.username.strip().lower()
    if len(username) < 3 or len(payload.password) < 8 or not payload.name.strip():
        raise ValueError("请填写有效用户名、姓名和至少 8 位密码")
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    application_id = f"application_{uuid.uuid4().hex[:12]}"
    try:
        with connect() as connection:
            connection.execute(
                """
                INSERT INTO users(
                    id, username, password_hash, name, role, organization, organization_id,
                    email, phone, account_status, department, title
                ) VALUES (?, ?, ?, ?, 'teacher', ?, ?, ?, ?, 'pending', ?, ?)
                """,
                (
                    user_id, username, hash_password(payload.password), payload.name.strip(),
                    payload.organization.strip() or "金扬智能示范学校", ORG_ID,
                    payload.email.strip(), payload.phone.strip(), payload.department.strip(), payload.title.strip(),
                ),
            )
            connection.execute(
                """
                INSERT INTO teacher_applications(id, user_id, department, title, status)
                VALUES (?, ?, ?, ?, 'pending')
                """,
                (application_id, user_id, payload.department.strip(), payload.title.strip()),
            )
    except Exception as exc:
        if "UNIQUE" in str(exc).upper():
            raise ValueError("用户名已存在") from exc
        raise
    return get_user(user_id)


def get_user(user_id: str) -> User:
    with connect() as connection:
        row = connection.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not row:
        raise KeyError("用户不存在")
    return _user_from_row(row)


def update_profile(user: User, payload: ProfileUpdate) -> User:
    if not payload.name.strip():
        raise ValueError("姓名不能为空")
    with connect() as connection:
        connection.execute(
            "UPDATE users SET name = ?, email = ?, phone = ?, title = ? WHERE id = ?",
            (payload.name.strip(), payload.email.strip(), payload.phone.strip(), payload.title.strip(), user.id),
        )
        _audit(connection, user, "profile.update", "user", user.id, {})
    return get_user(user.id)


def update_password(user: User, payload: PasswordUpdate) -> None:
    if len(payload.new_password) < 8:
        raise ValueError("新密码至少需要 8 个字符")
    with connect() as connection:
        row = connection.execute("SELECT password_hash FROM users WHERE id = ?", (user.id,)).fetchone()
        if not row or not verify_password(payload.current_password, row["password_hash"]):
            raise ValueError("当前密码错误")
        connection.execute(
            "UPDATE users SET password_hash = ?, must_change_password = 0 WHERE id = ?",
            (hash_password(payload.new_password), user.id),
        )
        _audit(connection, user, "password.change", "user", user.id, {})


def save_avatar(filename: str, raw: bytes, user: User) -> User:
    suffix = Path(filename).suffix.lower()
    if suffix not in (".png", ".jpg", ".jpeg", ".webp"):
        raise ValueError("头像仅支持 PNG、JPG 或 WEBP")
    directory = UPLOAD_DIR / "avatars"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{user.id}{suffix}"
    path.write_bytes(raw)
    avatar_url = f"/api/files/avatars/{path.name}"
    with connect() as connection:
        connection.execute("UPDATE users SET avatar_url = ? WHERE id = ?", (avatar_url, user.id))
    return get_user(user.id)


def list_classrooms(user: User) -> list[Classroom]:
    where = ""
    parameters: tuple[str, ...] = ()
    if user.role == "teacher":
        where = "WHERE EXISTS (SELECT 1 FROM class_teachers ct WHERE ct.classroom_id = c.id AND ct.teacher_id = ?)"
        parameters = (user.id,)
    elif user.role == "student":
        where = "WHERE c.status = 'active' AND EXISTS (SELECT 1 FROM enrollments e WHERE e.classroom_id = c.id AND e.student_id = ? AND e.status = 'active')"
        parameters = (user.id,)
    query = f"""
        SELECT c.*, co.name AS course_name, u.name AS primary_teacher_name,
          (SELECT COUNT(*) FROM class_teachers ct WHERE ct.classroom_id = c.id) AS teacher_count,
          (SELECT COUNT(*) FROM enrollments e WHERE e.classroom_id = c.id) AS student_count,
          (SELECT COUNT(*) FROM enrollments e JOIN users su ON su.id = e.student_id
            WHERE e.classroom_id = c.id AND e.status = 'active' AND su.account_status = 'active') AS active_student_count,
          COALESCE((SELECT AVG(progress_rate) FROM (
            SELECT e.student_id, 100.0 * SUM(COALESCE(p.mastered, 0)) / MAX(1, (SELECT COUNT(*) FROM nodes n WHERE n.course_id = c.course_id)) AS progress_rate
            FROM enrollments e LEFT JOIN class_learning_progress p ON p.classroom_id = e.classroom_id AND p.user_id = e.student_id
            WHERE e.classroom_id = c.id AND e.status = 'active' GROUP BY e.student_id
          )), 0) AS average_progress
        FROM classrooms c JOIN courses co ON co.id = c.course_id
        JOIN users u ON u.id = c.primary_teacher_id {where}
        ORDER BY c.updated_at DESC
    """
    with connect() as connection:
        rows = connection.execute(query, parameters).fetchall()
    return [_classroom_from_row(row) for row in rows]


def get_classroom(classroom_id: str, user: User) -> Classroom:
    found = [item for item in list_classrooms(user) if item.id == classroom_id]
    if not found:
        raise KeyError("班级不存在或无权访问")
    return found[0]


def create_classroom(payload: ClassroomCreate, user: User) -> Classroom:
    if user.role not in ("teacher", "admin"):
        raise PermissionError("当前角色不能创建班级")
    with connect() as connection:
        course = connection.execute("SELECT owner_id FROM courses WHERE id = ?", (payload.course_id,)).fetchone()
        if not course:
            raise KeyError("课程不存在")
        if user.role == "teacher" and course["owner_id"] != user.id:
            raise PermissionError("只能为自己创建的课程开班")
        class_id = f"class_{uuid.uuid4().hex[:10]}"
        code = _new_join_code(connection)
        connection.execute(
            """
            INSERT INTO classrooms(id, organization_id, course_id, name, semester, join_code, status, primary_teacher_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (class_id, ORG_ID, payload.course_id, payload.name.strip(), payload.semester.strip(), code, payload.status, user.id),
        )
        connection.execute(
            "INSERT INTO class_teachers(classroom_id, teacher_id, role, can_edit_course) VALUES (?, ?, 'primary', 1)",
            (class_id, user.id),
        )
        _audit(connection, user, "class.create", "classroom", class_id, payload.model_dump())
    return get_classroom(class_id, user)


def update_classroom(classroom_id: str, payload: ClassroomUpdate, user: User) -> Classroom:
    _assert_class_manage(classroom_id, user)
    with connect() as connection:
        connection.execute(
            "UPDATE classrooms SET name = ?, semester = ?, join_enabled = ?, status = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (payload.name.strip(), payload.semester.strip(), int(payload.join_enabled), payload.status, classroom_id),
        )
        _audit(connection, user, "class.update", "classroom", classroom_id, payload.model_dump())
    return get_classroom(classroom_id, user)


def reset_join_code(classroom_id: str, user: User) -> Classroom:
    _assert_class_manage(classroom_id, user)
    with connect() as connection:
        code = _new_join_code(connection)
        connection.execute("UPDATE classrooms SET join_code = ?, join_enabled = 1 WHERE id = ?", (code, classroom_id))
        _audit(connection, user, "class.join_code.reset", "classroom", classroom_id, {})
    return get_classroom(classroom_id, user)


def join_class(join_code: str, user: User) -> Classroom:
    if user.role != "student":
        raise PermissionError("仅学生可以通过班级码加入")
    with connect() as connection:
        row = connection.execute(
            "SELECT id FROM classrooms WHERE join_code = ? AND join_enabled = 1 AND status = 'active'",
            (join_code.strip().upper(),),
        ).fetchone()
        if not row:
            raise ValueError("班级码无效或班级已停止加入")
        connection.execute(
            """
            INSERT INTO enrollments(classroom_id, student_id, status) VALUES (?, ?, 'active')
            ON CONFLICT(classroom_id, student_id) DO UPDATE SET status = 'active'
            """,
            (row["id"], user.id),
        )
        _audit(connection, user, "class.join", "classroom", row["id"], {})
    return get_classroom(row["id"], user)


def list_class_members(classroom_id: str, user: User) -> list[User]:
    _assert_class_manage(classroom_id, user)
    with connect() as connection:
        rows = connection.execute(
            """
            SELECT u.* FROM enrollments e JOIN users u ON u.id = e.student_id
            WHERE e.classroom_id = ? AND e.status = 'active' ORDER BY u.student_no, u.name
            """,
            (classroom_id,),
        ).fetchall()
    return [_user_from_row(row) for row in rows]


def add_student_to_class(classroom_id: str, student_id: str, user: User) -> None:
    _assert_class_manage(classroom_id, user)
    with connect() as connection:
        target = connection.execute("SELECT role FROM users WHERE id = ?", (student_id,)).fetchone()
        if not target or target["role"] != "student":
            raise ValueError("目标用户不是学生")
        connection.execute(
            "INSERT INTO enrollments(classroom_id, student_id, status) VALUES (?, ?, 'active') ON CONFLICT(classroom_id, student_id) DO UPDATE SET status = 'active'",
            (classroom_id, student_id),
        )
        _audit(connection, user, "class.student.add", "classroom", classroom_id, {"student_id": student_id})


def assign_teacher(classroom_id: str, teacher_id: str, can_edit_course: bool, user: User) -> None:
    _assert_class_manage(classroom_id, user)
    with connect() as connection:
        target = connection.execute(
            "SELECT role, account_status FROM users WHERE id = ?", (teacher_id,)
        ).fetchone()
        if not target or target["role"] != "teacher" or target["account_status"] != "active":
            raise ValueError("只能添加已审核通过的教师")
        connection.execute(
            """
            INSERT INTO class_teachers(classroom_id, teacher_id, role, can_edit_course)
            VALUES (?, ?, 'collaborator', ?)
            ON CONFLICT(classroom_id, teacher_id) DO UPDATE SET can_edit_course = excluded.can_edit_course
            """,
            (classroom_id, teacher_id, int(can_edit_course)),
        )
        _audit(connection, user, "class.teacher.assign", "classroom", classroom_id, {"teacher_id": teacher_id})


def reassign_primary_teacher(classroom_id: str, teacher_id: str, user: User) -> Classroom:
    _require_admin(user)
    with connect() as connection:
        target = connection.execute(
            "SELECT role, account_status FROM users WHERE id = ?", (teacher_id,)
        ).fetchone()
        if not target or target["role"] != "teacher" or target["account_status"] != "active":
            raise ValueError("主教师必须是已审核通过的教师")
        connection.execute("UPDATE class_teachers SET role = 'collaborator' WHERE classroom_id = ?", (classroom_id,))
        connection.execute(
            """
            INSERT INTO class_teachers(classroom_id, teacher_id, role, can_edit_course)
            VALUES (?, ?, 'primary', 1)
            ON CONFLICT(classroom_id, teacher_id)
            DO UPDATE SET role = 'primary', can_edit_course = 1
            """,
            (classroom_id, teacher_id),
        )
        cursor = connection.execute(
            "UPDATE classrooms SET primary_teacher_id = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (teacher_id, classroom_id),
        )
        if cursor.rowcount == 0:
            raise KeyError("班级不存在")
        _audit(connection, user, "class.primary_teacher.update", "classroom", classroom_id, {"teacher_id": teacher_id})
    return get_classroom(classroom_id, user)


def admin_dashboard(user: User) -> DashboardStats:
    _require_admin(user)
    with connect() as connection:
        values = {
            "users": connection.execute("SELECT COUNT(*) FROM users").fetchone()[0],
            "teachers": connection.execute("SELECT COUNT(*) FROM users WHERE role = 'teacher'").fetchone()[0],
            "students": connection.execute("SELECT COUNT(*) FROM users WHERE role = 'student'").fetchone()[0],
            "pending_teachers": connection.execute("SELECT COUNT(*) FROM users WHERE role = 'teacher' AND account_status = 'pending'").fetchone()[0],
            "courses": connection.execute("SELECT COUNT(*) FROM courses").fetchone()[0],
            "classrooms": connection.execute("SELECT COUNT(*) FROM classrooms").fetchone()[0],
            "open_tickets": connection.execute("SELECT COUNT(*) FROM tickets WHERE status NOT IN ('resolved', 'closed')").fetchone()[0],
        }
    return DashboardStats(**values)


def admin_list_users(user: User, role: str = "", status: str = "", query: str = "") -> list[User]:
    _require_admin(user)
    clauses: list[str] = []
    params: list[str] = []
    if role:
        clauses.append("role = ?")
        params.append(role)
    if status:
        clauses.append("account_status = ?")
        params.append(status)
    if query:
        clauses.append("(name LIKE ? OR username LIKE ? OR student_no LIKE ?)")
        params.extend([f"%{query}%"] * 3)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    with connect() as connection:
        rows = connection.execute(f"SELECT * FROM users {where} ORDER BY created_at DESC", params).fetchall()
    return [_user_from_row(row) for row in rows]


def admin_update_user(user_id: str, payload: AdminUserUpdate, admin: User) -> User:
    _require_admin(admin)
    with connect() as connection:
        cursor = connection.execute(
            """
            UPDATE users SET name = ?, email = ?, phone = ?, account_status = ?, is_active = ?,
              department = ?, title = ?, student_no = ?, rejection_reason = ? WHERE id = ?
            """,
            (
                payload.name.strip(), payload.email.strip(), payload.phone.strip(), payload.account_status,
                int(payload.account_status != "disabled"), payload.department.strip(), payload.title.strip(),
                payload.student_no.strip(), "" if payload.account_status != "rejected" else "管理员驳回", user_id,
            ),
        )
        if cursor.rowcount == 0:
            raise KeyError("用户不存在")
        _audit(connection, admin, "admin.user.update", "user", user_id, payload.model_dump())
    return get_user(user_id)


def review_teacher(user_id: str, payload: TeacherReview, admin: User) -> User:
    _require_admin(admin)
    status = "active" if payload.action == "approve" else "rejected"
    application_status = "approved" if payload.action == "approve" else "rejected"
    if payload.action == "reject" and not payload.reason.strip():
        raise ValueError("驳回时必须填写原因")
    with connect() as connection:
        target = connection.execute("SELECT role FROM users WHERE id = ?", (user_id,)).fetchone()
        if not target or target["role"] != "teacher":
            raise KeyError("教师申请不存在")
        connection.execute(
            "UPDATE users SET account_status = ?, rejection_reason = ? WHERE id = ?",
            (status, payload.reason.strip(), user_id),
        )
        connection.execute(
            """
            UPDATE teacher_applications SET status = ?, reason = ?, reviewed_by = ?,
              reviewed_at = CURRENT_TIMESTAMP, updated_at = CURRENT_TIMESTAMP WHERE user_id = ?
            """,
            (application_status, payload.reason.strip(), admin.id, user_id),
        )
        _notify(connection, user_id, "教师申请审核结果", "申请已通过，可以开始开课。" if status == "active" else f"申请被驳回：{payload.reason}", "/profile")
        _audit(connection, admin, f"teacher.{payload.action}", "user", user_id, {"reason": payload.reason})
    return get_user(user_id)


def preview_student_import(classroom_id: str, filename: str, raw: bytes, user: User) -> ImportPreview:
    _assert_class_manage(classroom_id, user)
    rows = _read_import_rows(filename, raw)
    normalized: list[dict] = []
    valid = 0
    for index, row in enumerate(rows, start=2):
        student_no = str(row.get("学号") or row.get("student_no") or "").strip()
        name = str(row.get("姓名") or row.get("name") or "").strip()
        errors: list[str] = []
        if not student_no:
            errors.append("缺少学号")
        if not name:
            errors.append("缺少姓名")
        item = {
            "row": index, "student_no": student_no, "name": name,
            "email": str(row.get("邮箱") or row.get("email") or "").strip(),
            "phone": str(row.get("手机号") or row.get("phone") or "").strip(),
            "errors": errors,
        }
        if not errors:
            valid += 1
        normalized.append(item)
    job_id = f"import_{uuid.uuid4().hex[:12]}"
    with connect() as connection:
        connection.execute(
            "INSERT INTO import_jobs(id, classroom_id, requested_by, payload_json) VALUES (?, ?, ?, ?)",
            (job_id, classroom_id, user.id, json.dumps(normalized, ensure_ascii=False)),
        )
    return ImportPreview(job_id=job_id, total=len(normalized), valid=valid, invalid=len(normalized) - valid, rows=normalized)


def commit_student_import(job_id: str, user: User) -> ImportCommitResult:
    with connect() as connection:
        job = connection.execute("SELECT * FROM import_jobs WHERE id = ?", (job_id,)).fetchone()
        if not job:
            raise KeyError("导入任务不存在")
        _assert_class_manage(job["classroom_id"], user)
        if job["status"] != "preview":
            raise ValueError("该导入任务已经提交")
        rows = json.loads(job["payload_json"])
        created = updated = enrolled = 0
        credentials: list[dict[str, str]] = []
        for item in rows:
            if item["errors"]:
                continue
            existing = connection.execute("SELECT id FROM users WHERE username = ?", (item["student_no"],)).fetchone()
            if existing:
                student_id = existing["id"]
                connection.execute(
                    "UPDATE users SET name = ?, email = ?, phone = ?, student_no = ? WHERE id = ? AND role = 'student'",
                    (item["name"], item["email"], item["phone"], item["student_no"], student_id),
                )
                updated += 1
            else:
                student_id = f"user_{uuid.uuid4().hex[:12]}"
                password = _temporary_password()
                connection.execute(
                    """
                    INSERT INTO users(
                      id, username, password_hash, name, role, organization, organization_id,
                      email, phone, student_no, account_status, must_change_password
                    ) VALUES (?, ?, ?, ?, 'student', '金扬智能示范学校', ?, ?, ?, ?, 'active', 1)
                    """,
                    (student_id, item["student_no"], hash_password(password), item["name"], ORG_ID, item["email"], item["phone"], item["student_no"]),
                )
                credentials.append({"student_no": item["student_no"], "name": item["name"], "temporary_password": password})
                created += 1
            before = connection.total_changes
            connection.execute(
                "INSERT INTO enrollments(classroom_id, student_id, status) VALUES (?, ?, 'active') ON CONFLICT(classroom_id, student_id) DO UPDATE SET status = 'active'",
                (job["classroom_id"], student_id),
            )
            if connection.total_changes > before:
                enrolled += 1
        connection.execute("UPDATE import_jobs SET status = 'committed', committed_at = CURRENT_TIMESTAMP WHERE id = ?", (job_id,))
        _audit(connection, user, "students.import", "classroom", job["classroom_id"], {"created": created, "updated": updated})
    return ImportCommitResult(created=created, updated=updated, enrolled=enrolled, credentials=credentials)


def create_ticket(payload: TicketCreate, user: User) -> Ticket:
    if len(payload.subject.strip()) < 3 or len(payload.description.strip()) < 5:
        raise ValueError("请完整填写反馈标题和问题描述")
    ticket_id = f"ticket_{uuid.uuid4().hex[:12]}"
    with connect() as connection:
        connection.execute(
            """
            INSERT INTO tickets(id, organization_id, creator_id, category, severity, subject, description)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (ticket_id, ORG_ID, user.id, payload.category, payload.severity, payload.subject.strip(), payload.description.strip()),
        )
        admins = connection.execute("SELECT id FROM users WHERE role = 'admin' AND account_status = 'active'").fetchall()
        for admin in admins:
            _notify(connection, admin["id"], "收到新反馈", payload.subject.strip(), f"/tickets/{ticket_id}")
        _audit(connection, user, "ticket.create", "ticket", ticket_id, {"severity": payload.severity})
    return get_ticket(ticket_id, user)


def list_tickets(user: User) -> list[Ticket]:
    where = "" if user.role == "admin" else "WHERE t.creator_id = ?"
    params = () if user.role == "admin" else (user.id,)
    with connect() as connection:
        rows = connection.execute(
            f"""
            SELECT t.*, creator.name AS creator_name, COALESCE(admin.name, '') AS assigned_admin_name
            FROM tickets t JOIN users creator ON creator.id = t.creator_id
            LEFT JOIN users admin ON admin.id = t.assigned_admin_id {where}
            ORDER BY CASE t.severity WHEN 'urgent' THEN 0 WHEN 'high' THEN 1 ELSE 2 END, t.updated_at DESC
            """,
            params,
        ).fetchall()
    return [_ticket_from_row(row) for row in rows]


def get_ticket(ticket_id: str, user: User) -> Ticket:
    with connect() as connection:
        row = connection.execute(
            """
            SELECT t.*, creator.name AS creator_name, COALESCE(admin.name, '') AS assigned_admin_name
            FROM tickets t JOIN users creator ON creator.id = t.creator_id
            LEFT JOIN users admin ON admin.id = t.assigned_admin_id WHERE t.id = ?
            """,
            (ticket_id,),
        ).fetchone()
        if not row or (user.role != "admin" and row["creator_id"] != user.id):
            raise KeyError("工单不存在或无权查看")
        messages = [dict(item) for item in connection.execute(
            """
            SELECT m.id, m.content, m.created_at, u.id AS author_id, u.name AS author_name, u.role AS author_role
            FROM ticket_messages m JOIN users u ON u.id = m.author_id
            WHERE m.ticket_id = ? ORDER BY m.created_at
            """,
            (ticket_id,),
        ).fetchall()]
        attachments = [dict(item) for item in connection.execute(
            "SELECT id, filename, size, created_at FROM ticket_attachments WHERE ticket_id = ? ORDER BY created_at",
            (ticket_id,),
        ).fetchall()]
        for item in attachments:
            item["url"] = f"/api/tickets/{ticket_id}/attachments/{item['id']}"
    ticket = _ticket_from_row(row)
    ticket.messages = messages
    ticket.attachments = attachments
    return ticket


def save_ticket_attachment(ticket_id: str, filename: str, raw: bytes, user: User) -> Ticket:
    get_ticket(ticket_id, user)
    suffix = Path(filename).suffix.lower()
    if suffix not in (".png", ".jpg", ".jpeg", ".webp", ".pdf"):
        raise ValueError("工单附件仅支持图片或 PDF")
    attachment_id = f"attachment_{uuid.uuid4().hex[:12]}"
    safe_name = Path(filename).name
    directory = UPLOAD_DIR / "tickets" / ticket_id
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{attachment_id}_{safe_name}"
    path.write_bytes(raw)
    with connect() as connection:
        connection.execute(
            """
            INSERT INTO ticket_attachments(id, ticket_id, filename, storage_path, size)
            VALUES (?, ?, ?, ?, ?)
            """,
            (attachment_id, ticket_id, safe_name, str(path), len(raw)),
        )
        _audit(connection, user, "ticket.attachment.add", "ticket", ticket_id, {"filename": safe_name})
    return get_ticket(ticket_id, user)


def ticket_attachment_path(ticket_id: str, attachment_id: str, user: User) -> tuple[Path, str]:
    get_ticket(ticket_id, user)
    with connect() as connection:
        row = connection.execute(
            "SELECT storage_path, filename FROM ticket_attachments WHERE id = ? AND ticket_id = ?",
            (attachment_id, ticket_id),
        ).fetchone()
    if not row:
        raise KeyError("附件不存在")
    return Path(row["storage_path"]), row["filename"]


def list_notifications(user: User) -> list[dict]:
    with connect() as connection:
        return [dict(row) for row in connection.execute(
            "SELECT * FROM notifications WHERE user_id = ? ORDER BY created_at DESC LIMIT 100", (user.id,)
        ).fetchall()]


def mark_notification_read(notification_id: str, user: User) -> None:
    with connect() as connection:
        cursor = connection.execute(
            "UPDATE notifications SET is_read = 1 WHERE id = ? AND user_id = ?", (notification_id, user.id)
        )
        if cursor.rowcount == 0:
            raise KeyError("通知不存在")


def reply_ticket(ticket_id: str, payload: TicketReplyCreate, user: User) -> Ticket:
    ticket = get_ticket(ticket_id, user)
    if not payload.content.strip():
        raise ValueError("回复内容不能为空")
    with connect() as connection:
        connection.execute(
            "INSERT INTO ticket_messages(id, ticket_id, author_id, content) VALUES (?, ?, ?, ?)",
            (f"message_{uuid.uuid4().hex[:12]}", ticket_id, user.id, payload.content.strip()),
        )
        next_status = "waiting_user" if user.role == "admin" else "in_progress"
        connection.execute("UPDATE tickets SET status = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (next_status, ticket_id))
        target = ticket.creator_id if user.role == "admin" else ticket.assigned_admin_id
        if target:
            _notify(connection, target, "工单收到新回复", ticket.subject, f"/tickets/{ticket_id}")
    return get_ticket(ticket_id, user)


def update_ticket(ticket_id: str, payload: TicketUpdate, admin: User) -> Ticket:
    _require_admin(admin)
    with connect() as connection:
        connection.execute(
            "UPDATE tickets SET status = ?, assigned_admin_id = NULLIF(?, ''), updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (payload.status, payload.assigned_admin_id, ticket_id),
        )
        row = connection.execute("SELECT creator_id, subject FROM tickets WHERE id = ?", (ticket_id,)).fetchone()
        if not row:
            raise KeyError("工单不存在")
        _notify(connection, row["creator_id"], "工单状态已更新", f"{row['subject']}：{payload.status}", f"/tickets/{ticket_id}")
        _audit(connection, admin, "ticket.update", "ticket", ticket_id, payload.model_dump())
    return get_ticket(ticket_id, admin)


def list_audit_logs(user: User, limit: int = 100) -> list[dict]:
    _require_admin(user)
    with connect() as connection:
        return [dict(row) for row in connection.execute(
            """
            SELECT a.*, u.name AS actor_name FROM audit_logs a JOIN users u ON u.id = a.actor_id
            ORDER BY a.created_at DESC LIMIT ?
            """,
            (min(max(limit, 1), 500),),
        ).fetchall()]


def get_class_graph(classroom_id: str, user: User) -> KnowledgeGraph:
    classroom = get_classroom(classroom_id, user)
    from app import storage
    graph = storage.get_graph(classroom.course_id, user if user.role != "student" else User(**user.model_dump()))
    if user.role == "student":
        mastered = set(class_mastered_node_ids(classroom_id, user))
        graph.nodes = [node.model_copy(update={"mastered": node.id in mastered}) for node in graph.nodes]
    return graph


def set_class_progress(classroom_id: str, node_id: str, mastered: bool, user: User) -> KnowledgeNode:
    if user.role != "student":
        raise PermissionError("仅学生可以记录学习进度")
    classroom = get_classroom(classroom_id, user)
    with connect() as connection:
        node = connection.execute("SELECT * FROM nodes WHERE id = ? AND course_id = ?", (node_id, classroom.course_id)).fetchone()
        if not node:
            raise KeyError("知识点不存在")
        connection.execute(
            """
            INSERT INTO class_learning_progress(user_id, classroom_id, node_id, mastered, updated_at)
            VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, classroom_id, node_id)
            DO UPDATE SET mastered = excluded.mastered, updated_at = CURRENT_TIMESTAMP
            """,
            (user.id, classroom_id, node_id, int(mastered)),
        )
    return KnowledgeNode(
        id=node["id"], name=node["name"], type=node["type"], definition=node["definition"],
        example=node["example"], resources=json.loads(node["resources_json"] or "[]"), mastered=mastered,
    )


def class_mastered_node_ids(classroom_id: str, user: User) -> list[str]:
    get_classroom(classroom_id, user)
    with connect() as connection:
        rows = connection.execute(
            "SELECT node_id FROM class_learning_progress WHERE user_id = ? AND classroom_id = ? AND mastered = 1",
            (user.id, classroom_id),
        ).fetchall()
    return [row["node_id"] for row in rows]


def diagnosis(classroom_id: str, user: User) -> DiagnosisResult:
    from app import graph_lifecycle
    from app.services import deepseek
    from app.services.recommender import recommend_path_for_course

    graph = get_class_graph(classroom_id, user)
    mastered = [node for node in graph.nodes if node.mastered]
    weak = [node for node in graph.nodes if not node.mastered][:5]
    rate = round(100 * len(mastered) / max(1, len(graph.nodes)), 1)
    suggestions = [f"先学习“{node.name}”，再沿前置关系继续。" for node in weak[:3]]
    if not suggestions:
        suggestions = ["本课程知识点已全部标记掌握，可进入综合练习。"]
    classroom = get_classroom(classroom_id, user)
    path = recommend_path_for_course(classroom.course_id, graph, [node.id for node in mastered])
    query = " ".join(node.name for node in weak[:5])
    evidence, _ = graph_lifecycle.retrieve_evidence(classroom.course_id, query, user) if query else ([], [])
    with connect() as connection:
        course = connection.execute("SELECT name FROM courses WHERE id = ?", (classroom.course_id,)).fetchone()
    try:
        analysis = deepseek.learning_analysis(
            course["name"] if course else classroom.course_name,
            rate,
            weak,
            [item.node for item in path.recommendations],
            evidence,
        )
        mode = "deepseek-evidence" if deepseek.configured() else "offline-rule"
    except ValueError:
        analysis = "AI 服务暂不可用，当前建议仍由掌握状态和前置关系规则生成。"
        mode = "offline-fallback"
    return DiagnosisResult(
        mastery_rate=rate, mastered_count=len(mastered), total_count=len(graph.nodes), weak_nodes=weak,
        suggestions=suggestions, ai_analysis=analysis, mode=mode, evidence=evidence[:4],
    )


def generate_exercises(classroom_id: str, user: User) -> list[ExerciseResult]:
    from app import graph_lifecycle
    from app.services import deepseek

    graph = get_class_graph(classroom_id, user)
    weak_nodes = [node for node in graph.nodes if not node.mastered][:5]
    if not weak_nodes:
        weak_nodes = graph.nodes[:3]
    classroom = get_classroom(classroom_id, user)
    query = " ".join(node.name for node in weak_nodes)
    evidence, _ = graph_lifecycle.retrieve_evidence(classroom.course_id, query, user) if query else ([], [])
    try:
        return deepseek.exercises(classroom.course_name, weak_nodes, evidence)
    except ValueError:
        return deepseek.offline_exercises(weak_nodes[:3], evidence)


def _graph_dict(connection, course_id: str) -> dict:
    nodes = [
        {
            "id": row["id"], "name": row["name"], "type": row["type"], "definition": row["definition"],
            "example": row["example"], "resources": json.loads(row["resources_json"] or "[]"), "mastered": False,
        }
        for row in connection.execute("SELECT * FROM nodes WHERE course_id = ? ORDER BY created_at, id", (course_id,)).fetchall()
    ]
    edges = [dict(row) for row in connection.execute(
        "SELECT id, source, target, relation, label FROM edges WHERE course_id = ? ORDER BY created_at, id", (course_id,)
    ).fetchall()]
    return {"nodes": nodes, "edges": edges}


def _classroom_from_row(row) -> Classroom:
    return Classroom(
        id=row["id"], course_id=row["course_id"], course_name=row["course_name"], name=row["name"],
        semester=row["semester"], join_code=row["join_code"], join_enabled=bool(row["join_enabled"]),
        status=row["status"], primary_teacher_id=row["primary_teacher_id"],
        primary_teacher_name=row["primary_teacher_name"], teacher_count=row["teacher_count"],
        student_count=row["student_count"], active_student_count=row["active_student_count"],
        average_progress=round(float(row["average_progress"] or 0), 1),
    )


def _user_from_row(row) -> User:
    return User(
        id=row["id"], name=row["name"], role=row["role"], organization=row["organization"],
        username=row["username"], email=row["email"], phone=row["phone"], avatar_url=row["avatar_url"],
        account_status=row["account_status"], must_change_password=bool(row["must_change_password"]),
        student_no=row["student_no"], department=row["department"], title=row["title"],
        rejection_reason=row["rejection_reason"],
    )


def _ticket_from_row(row) -> Ticket:
    return Ticket(
        id=row["id"], creator_id=row["creator_id"], creator_name=row["creator_name"], category=row["category"],
        severity=row["severity"], subject=row["subject"], description=row["description"], status=row["status"],
        assigned_admin_id=row["assigned_admin_id"] or "", assigned_admin_name=row["assigned_admin_name"],
        created_at=row["created_at"], updated_at=row["updated_at"],
    )


def _assert_class_manage(classroom_id: str, user: User) -> None:
    if user.role == "admin":
        return
    if user.role != "teacher" or user.account_status != "active":
        raise PermissionError("无权管理该班级")
    with connect() as connection:
        row = connection.execute(
            "SELECT 1 FROM class_teachers WHERE classroom_id = ? AND teacher_id = ?", (classroom_id, user.id)
        ).fetchone()
    if not row:
        raise PermissionError("无权管理该班级")


def _require_admin(user: User) -> None:
    if user.role != "admin":
        raise PermissionError("仅管理员可以执行此操作")


def _new_join_code(connection) -> str:
    for _ in range(20):
        code = "".join(secrets.choice(string.ascii_uppercase + string.digits) for _ in range(6))
        if not connection.execute("SELECT 1 FROM classrooms WHERE join_code = ?", (code,)).fetchone():
            return code
    raise RuntimeError("无法生成班级码")


def _temporary_password() -> str:
    alphabet = string.ascii_letters + string.digits
    return "Cg!" + "".join(secrets.choice(alphabet) for _ in range(9))


def _read_import_rows(filename: str, raw: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = raw.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("服务器未安装 XLSX 解析组件，请改用 CSV") from exc
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:]]
    raise ValueError("仅支持 CSV 或 XLSX 导入")


def _notify(connection, user_id: str, title: str, content: str, link: str) -> None:
    connection.execute(
        "INSERT INTO notifications(id, user_id, title, content, link) VALUES (?, ?, ?, ?, ?)",
        (f"notification_{uuid.uuid4().hex[:12]}", user_id, title, content, link),
    )


def _audit(connection, user: User, action: str, target_type: str, target_id: str, detail: dict) -> None:
    connection.execute(
        """
        INSERT INTO audit_logs(id, organization_id, actor_id, action, target_type, target_id, detail_json)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (f"audit_{uuid.uuid4().hex[:12]}", ORG_ID, user.id, action, target_type, target_id, json.dumps(detail, ensure_ascii=False)),
    )
